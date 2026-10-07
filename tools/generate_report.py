from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidencia" / "salidas"
DOCX_PATH = ROOT / "Informe_Avance_2_DFD.docx"
PDF_FILENAME = "Informe_Avance_2_DFD_final.pdf"
PDF_PATH = ROOT / PDF_FILENAME


LEXICAL_HEADERS = ["componente léxico", "Expresión regular", "valor del atributo"]
LEXICAL_ROWS = [
    ("reservada", "SISTEMA | FIN_SISTEMA | NIVEL | entidad | proceso | almacen | flujo | conectar | con_dato | si | sino | mientras | entero | booleano | true | false | imprimir | validar_modelo | contar_nodos", "token reservado"),
    ("oprel", "=", "IGUAL"),
    ("oprel", "<", "MEN"),
    ("oprel", "<=", "MEI"),
    ("oprel", ">", "MAY"),
    ("oprel", ">=", "MAI"),
    ("oprel", "!=", "DIF"),
    ("oprel", "==", "IGU"),
    ("oparit", "+", "SUM"),
    ("oparit", "-", "RES"),
    ("oparit", "*", "MUL"),
    ("oparit", "/", "DIV"),
    ("oplogico", "&&", "AND"),
    ("oplogico", "||", "OR"),
    ("flujo", "->", "FLECHA"),
    ("puntuación", "(", "PA"),
    ("puntuación", ")", "PC"),
    ("puntuación", "{", "LLA"),
    ("puntuación", "}", "LLC"),
    ("puntuación", ",", "COMA"),
    ("puntuación", ";", "PYC"),
    ("id", r"[a-zA-Z][a-zA-Z0-9_]*", "puntero a la tabla de símbolos (TS)"),
    ("const", r"[0-9]+ ('.' [0-9]+)?", "puntero a la tabla de símbolos (TS)"),
    ("cadena", r"'\"' ~[\"\\r\\n]* '\"'", "CADENA"),
    ("espacios", r"[ \\t\\r\\n]+", "skip"),
    ("error", ".", "ERROR"),
]


def read_output(name: str) -> str:
    raw = (EVIDENCE / name).read_bytes()
    if raw.startswith(b"\xff\xfe"):
        return raw.decode("utf-16").strip()
    return raw.decode("utf-8").strip()


def lines(name: str):
    return read_output(name).splitlines()


def token_count(name: str) -> int:
    return len(lines(name))


def setup_docx(document: Document):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for style_name, size, color in (("Title", 19, "1F4E79"), ("Heading 1", 14, "1F4E79"), ("Heading 2", 11.5, "2F5597")):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Informe de avance 2 - Lenguaje DFD - Taller 2026")


def add_code_docx(document: Document, text: str):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.2)


def add_table_docx(document: Document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()
    return table


def build_docx():
    document = Document()
    setup_docx(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Informe de avance 2\nDesarrollo del analizador léxico para un lenguaje DFD")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Lenguajes de Programación y Compiladores 2026 - UNSE FCEyT")
    document.add_paragraph()
    document.add_paragraph("Grupo 5: Gastón Coronel, Jerónimo Gili, Leonardo Noriega y Jorge Taboada Nuno.")
    document.add_paragraph("Datos que deben verificarse antes de entregar: legajo de Leonardo Noriega, legajo de Jorge Taboada Nuno, nombre definitivo del grupo y continuidad de los cuatro integrantes.")

    document.add_heading("1. Lenguaje y alcance", level=1)
    document.add_paragraph(
        "El lenguaje describe diagramas de flujo de datos (DFD) mediante instrucciones textuales. "
        "Un DFD muestra cómo circula información entre entidades externas, procesos y almacenes. "
        "En este avance se implementa únicamente el analizador léxico: lee caracteres y produce tokens "
        "con tipo, lexema, línea y columna. Todavía no se valida la estructura de las sentencias ni la "
        "validez semántica de las conexiones."
    )
    add_table_docx(document, ["Elemento", "Explicación"], [
        ("Entidad", "Participante externo que origina o recibe información, por ejemplo Cliente."),
        ("Proceso", "Actividad que transforma información, por ejemplo Gestionar."),
        ("Almacén", "Lugar lógico donde se conservan datos, por ejemplo una base de pedidos."),
        ("Flujo", "Conexión que transporta un dato entre elementos, por ejemplo Pedido."),
    ])
    document.add_paragraph("Programa principal utilizado como caso de prueba:")
    add_code_docx(document, (ROOT / "src/test/resources/programa_informe1.dfd").read_text(encoding="utf-8").strip())

    document.add_heading("2. Antecedentes del informe de avance 1", level=1)
    document.add_paragraph("El Informe 1 definió el lenguaje DFD, su vocabulario, las expresiones regulares, la gramática libre de contexto y una prueba manual. Esta sección conserva esa base; las correcciones se presentan en la sección siguiente.")
    document.add_paragraph("Vocabulario terminal de referencia:")
    add_code_docx(document, "SISTEMA, FIN_SISTEMA, NIVEL, entidad, proceso, almacen, flujo, conectar, con_dato,\nsi, sino, mientras, entero, booleano, true, false, imprimir, validar_modelo, contar_nodos,\nID, CADENA, NUMERO, ->, ==, !=, <=, >=, <, >, &&, ||, =, +, -, *, /, ;, {, }, (, )")
    document.add_paragraph("Expresiones regulares publicadas como base del Informe 1:")
    add_code_docx(document, "Identificadores: [a-zA-Z][a-zA-Z0-9]*\nNúmeros: -?(0|[1-9][0-9]*(\\.[0-9]+)?)\nCadenas: \"[^\"]*\"")
    document.add_paragraph("La GLC del Informe 1 establece la estructura general del programa y queda como referencia para el avance sintáctico:")
    add_code_docx(document, "Programa -> SISTEMA ID ; ListaNiveles FIN_SISTEMA\nDefNivel -> NIVEL NUMERO Bloque\nBloque -> { ListaSentencias }\nSentencia -> DeclaracionNodo | DeclaracionFlujo | DeclaracionVar | Condicional | Iterativa | LlamadaFuncion\nDeclaracionNodo -> TipoNodo ID = CADENA ;\nTipoNodo -> entidad | proceso | almacen\nDeclaracionFlujo -> flujo ID = conectar ( ID -> ID ) con_dato CADENA ;")
    document.add_paragraph("El lenguaje también contempla variables de tipo entero y booleano, estructuras si/sino y mientras, y funciones del dominio como imprimir, validar_modelo y contar_nodos. Estas construcciones pertenecen al análisis sintáctico y semántico posterior; en este avance solo se reconocen sus palabras reservadas y nombres.")
    document.add_paragraph("La regla original ExpresionLogica reunía comparaciones, conjunción y disyunción en una misma producción. En este avance no se reescribe el parser: solo se separan léxicamente sus operadores.")
    add_code_docx(document, "ExpresionLogica -> Expresion ExpLogR\nExpLogR -> == Expresion ExpLogR | != Expresion ExpLogR | < Expresion ExpLogR | > Expresion ExpLogR\n         | <= Expresion ExpLogR | >= Expresion ExpLogR | && Expresion ExpLogR | || Expresion ExpLogR | lambda")
    document.add_paragraph("La prueba manual del Informe 1 utiliza el mismo programa Ventas mostrado en la sección 1. La validación sintáctica de esa prueba corresponde al avance 3.")

    document.add_heading("3. Correcciones al informe de avance 1", level=1)
    document.add_paragraph("La devolución indicó que faltaba explicar el lenguaje, aclarar la expresión regular de cadenas y desagregar ExpresionLogica. Las correcciones léxicas aplicadas son:")
    corrections = [
        "El texto decía que un identificador admitía guion bajo, pero la expresión regular publicada no lo incluía. Se adopta [a-zA-Z][a-zA-Z0-9_]*.",
        "La expresión de cadenas se precisa como comilla inicial, cero o más caracteres que no sean comilla ni salto de línea y comilla final: CADENA : '\"' ~[\"\\r\\n]* '\"' ;. Se permiten cadenas vacías y no se admiten escapes.",
        "El vocabulario incluye almacen, pero no aparecía en todos los lugares de forma consistente. Se incorpora como palabra reservada ALMACEN.",
        "El signo menos se reconoce separado del NUMERO para que el avance sintáctico pueda distinguir resta de número negativo.",
        "Los operadores ==, !=, <, >, <=, >=, && y || tienen tokens separados. La desagregación de ExpresionLogica en niveles de comparación, conjunción y disyunción se desarrollará en el avance sintáctico.",
    ]
    for item in corrections:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("4. Componentes léxicos", level=1)
    add_table_docx(document, LEXICAL_HEADERS, LEXICAL_ROWS)

    document.add_heading("5. Implementación con ANTLR 4", level=1)
    document.add_paragraph("La gramática lexer DFDLexer.g4 contiene las reglas del vocabulario. ANTLR aplica la coincidencia más larga; cuando dos reglas tienen la misma longitud, la regla que aparece primero tiene prioridad. Por eso entidad es ENTIDAD, entidad1 es ID y -> es FLECHA.")
    add_code_docx(document, "ENTIDAD : 'entidad';\nFLECHA  : '->';\nNUMERO  : [0-9]+ ('.' [0-9]+)?;\nCADENA  : '\"' ~[\"\\r\\n]* '\"';\nID      : [a-zA-Z] [a-zA-Z0-9_]*;\nESPACIOS: [ \\t\\r\\n]+ -> skip;\nERROR   : .;")
    document.add_paragraph("El programa MostrarTokens.java recibe un archivo DFD y muestra cada token con el formato TIPO | LEXEMA | LINEA | COLUMNA. No crea todavía un árbol sintáctico.")
    document.add_paragraph("Ejecución utilizada:")
    add_code_docx(document, "mvn -q clean package\nmvn -q exec:java '-Dexec.mainClass=org.example.MostrarTokens' '-Dexec.args=src/test/resources/programa_informe1.dfd'")

    document.add_heading("6. Pruebas del analizador", level=1)
    document.add_paragraph(f"La ejecución del programa principal produjo {token_count('programa_informe1.txt')} tokens y no produjo tokens ERROR.")
    document.add_paragraph("Primeras y últimas líneas de la salida real:")
    main_lines = lines("programa_informe1.txt")
    add_code_docx(document, "\n".join(main_lines[:8] + ["...", *main_lines[-6:]]))
    add_table_docx(document, ["Entrada", "Resultado real resumido"], [
        ("reservadas_identificadores.dfd", "Las 19 palabras reservadas se reconocen como reservadas; SISTEMA1, entidad1, cliente_externo y bd_01 se reconocen como ID."),
        ("numeros.dfd", "0, 12 y 3.5 son NUMERO; -2 es MENOS seguido de NUMERO; puntos mal ubicados producen ERROR."),
        ("cadenas.dfd", "\"Pedido\", \"\" y \"Cliente principal\" son CADENA."),
        ("operadores.dfd", "->, ==, !=, <=, >=, && y || se reconocen como tokens individuales de dos caracteres."),
        ("invalidos.dfd", "@, #, $, y el guion bajo inicial producen ERROR; 1proceso se divide en NUMERO e ID."),
        ("cadena_sin_cierre.dfd", "La comilla produce ERROR y Pedido sin cierre se reconoce como tres ID. Es una limitación léxica documentada."),
    ])
    document.add_paragraph("Las salidas completas se conservan en evidencia/salidas/. Las pruebas verifican reconocimiento léxico, no corrección sintáctica ni semántica.")

    document.add_heading("7. Cronograma de la etapa", level=1)
    add_table_docx(document, ["Actividad", "Período", "Responsable", "Estado"], [
        ("Revisar devolución y fijar lenguaje DFD", "Hasta 22/09/2026", "Equipo", "Realizado"),
        ("Ajustar vocabulario y escribir lexer ANTLR", "22/09/2026", "Equipo", "Realizado"),
        ("Generar analizador y documentar pruebas", "22-23/09/2026", "Equipo", "Realizado con salidas guardadas"),
        ("Revisar datos del grupo, redacción y entrega", "23/09/2026", "Un integrante y equipo", "Verificación final pendiente"),
    ])

    document.add_heading("8. Trabajo previsto para el avance sintáctico", level=1)
    document.add_paragraph("En el avance 3 se reemplazará la regla monolítica ExpresionLogica por niveles separados para comparación, conjunción y disyunción, conservando la precedencia de && sobre || y utilizando los tokens ya generados en este avance. El parser no forma parte de la entrega actual.")

    document.add_heading("9. Archivos y verificaciones antes de entregar", level=1)
    document.add_paragraph("El archivo que corresponde presentar según el enunciado es el PDF. El DOCX se conserva como fuente editable y el proyecto con sus evidencias permite reproducir el analizador.")
    add_table_docx(document, ["Elemento", "Acción"], [
        (PDF_FILENAME, "Presentar en el aula virtual, después de completar los datos del grupo."),
        ("Informe_Avance_2_DFD.docx", "Conservar como editable; presentarlo solo si la cátedra lo solicita."),
        ("Proyecto Maven/ANTLR", "Conservar con pom.xml, DFDLexer.g4, MostrarTokens.java, pruebas y evidencia/salidas/."),
        ("Datos del grupo", "Completar legajos faltantes, confirmar nombres, grupo y continuidad de integrantes."),
    ])

    document.add_heading("Referencias", level=1)
    document.add_paragraph("Coronel, G. y Gili, J. (2026). Informe de avance 1: Taller 2026.")
    document.add_paragraph("Devolución de la entrega 1 - Coronel, Gili (2026).")
    document.add_paragraph("Cátedra de Lenguajes de Programación y Compiladores (2026). Enunciado del taller 2026. UNSE FCEyT.")
    document.add_paragraph("Las referencias legacy se consultaron solo para observar el formato histórico de trabajos de la cátedra; no se reutilizó un informe de otro lenguaje.")

    document.save(DOCX_PATH)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleBlue", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#1F4E79"), alignment=TA_CENTER, spaceAfter=10))
    styles.add(ParagraphStyle(name="HeadingBlue", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#1F4E79"), spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="BodyES", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.2, spaceAfter=6))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9, spaceAfter=3))
    styles.add(ParagraphStyle(name="CodeES", parent=styles["Code"], fontName="Courier", fontSize=7.2, leading=8.4, leftIndent=8, rightIndent=8, spaceBefore=3, spaceAfter=6))
    return styles


def p(text, style):
    return Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), style)


def pdf_table(data, widths, styles, header=True, font_size=7.2):
    converted = []
    for r, row in enumerate(data):
        converted.append([Paragraph(str(cell).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), styles["Small"]) for cell in row])
    table = Table(converted, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E8F5")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9FBAD0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_pdf():
    styles = pdf_styles()
    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    story = []
    story.append(Paragraph("Informe de avance 2", styles["TitleBlue"]))
    story.append(Paragraph("Desarrollo del analizador léxico para un lenguaje DFD", styles["TitleBlue"]))
    story.append(Paragraph("Lenguajes de Programación y Compiladores 2026 - UNSE FCEyT", styles["BodyES"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Grupo 5: Gastón Coronel, Jerónimo Gili, Leonardo Noriega y Jorge Taboada Nuno.", styles["BodyES"]))
    story.append(Paragraph("Verificar antes de entregar: legajos de Leonardo Noriega y Jorge Taboada Nuno, nombre definitivo del grupo y continuidad de los cuatro integrantes.", styles["BodyES"]))

    story.append(Paragraph("1. Lenguaje y alcance", styles["HeadingBlue"]))
    story.append(Paragraph("El lenguaje describe diagramas de flujo de datos (DFD) mediante instrucciones textuales. Una entidad es un participante externo; un proceso transforma información; un almacén conserva datos; y un flujo transporta información entre elementos. Este avance implementa únicamente el analizador léxico: produce tokens con tipo, lexema, línea y columna. No valida todavía la estructura ni la semántica de las sentencias.", styles["BodyES"]))
    story.append(pdf_table([["Elemento", "Explicación"], ["Entidad", "Participante externo que origina o recibe información."], ["Proceso", "Actividad que transforma información."], ["Almacén", "Lugar lógico donde se conservan datos."], ["Flujo", "Conexión que transporta un dato entre elementos."]], [3.2 * cm, 13.8 * cm], styles))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Caso principal del informe 1", styles["BodyES"]))
    story.append(Preformatted((ROOT / "src/test/resources/programa_informe1.dfd").read_text(encoding="utf-8").strip(), styles["CodeES"]))

    story.append(Paragraph("2. Antecedentes del informe de avance 1", styles["HeadingBlue"]))
    story.append(Paragraph("El Informe 1 definió el lenguaje DFD, su vocabulario, las expresiones regulares, la gramática libre de contexto y una prueba manual. Esta sección conserva esa base; las correcciones se presentan en la sección siguiente.", styles["BodyES"]))
    story.append(Paragraph("Vocabulario terminal de referencia:", styles["BodyES"]))
    story.append(Preformatted("SISTEMA, FIN_SISTEMA, NIVEL, entidad, proceso, almacen, flujo, conectar, con_dato,\nsi, sino, mientras, entero, booleano, true, false, imprimir, validar_modelo, contar_nodos,\nID, CADENA, NUMERO, ->, ==, !=, <=, >=, <, >, &&, ||, =, +, -, *, /, ;, {, }, (, )", styles["CodeES"]))
    story.append(Paragraph("Expresiones regulares publicadas como base del Informe 1:", styles["BodyES"]))
    story.append(Preformatted("Identificadores: [a-zA-Z][a-zA-Z0-9]*\nNúmeros: -?(0|[1-9][0-9]*(\\.[0-9]+)?)\nCadenas: \"[^\"]*\"", styles["CodeES"]))
    story.append(Paragraph("La GLC del Informe 1 establece la estructura general del programa y queda como referencia para el avance sintáctico:", styles["BodyES"]))
    story.append(Preformatted("Programa -> SISTEMA ID ; ListaNiveles FIN_SISTEMA\nDefNivel -> NIVEL NUMERO Bloque\nBloque -> { ListaSentencias }\nSentencia -> DeclaracionNodo | DeclaracionFlujo | DeclaracionVar | Condicional | Iterativa | LlamadaFuncion\nDeclaracionNodo -> TipoNodo ID = CADENA ;\nTipoNodo -> entidad | proceso | almacen\nDeclaracionFlujo -> flujo ID = conectar ( ID -> ID ) con_dato CADENA ;", styles["CodeES"]))
    story.append(Paragraph("El lenguaje también contempla variables de tipo entero y booleano, estructuras si/sino y mientras, y funciones del dominio como imprimir, validar_modelo y contar_nodos. Estas construcciones pertenecen al análisis sintáctico y semántico posterior; en este avance solo se reconocen sus palabras reservadas y nombres.", styles["BodyES"]))
    story.append(Paragraph("La regla original ExpresionLogica reunía comparaciones, conjunción y disyunción en una misma producción. En este avance no se reescribe el parser: solo se separan léxicamente sus operadores.", styles["BodyES"]))
    story.append(Preformatted("ExpresionLogica -> Expresion ExpLogR\nExpLogR -> == Expresion ExpLogR | != Expresion ExpLogR | < Expresion ExpLogR | > Expresion ExpLogR\n         | <= Expresion ExpLogR | >= Expresion ExpLogR | && Expresion ExpLogR | || Expresion ExpLogR | lambda", styles["CodeES"]))
    story.append(Paragraph("La prueba manual del Informe 1 utiliza el mismo programa Ventas mostrado en la sección 1. La validación sintáctica de esa prueba corresponde al avance 3.", styles["BodyES"]))

    story.append(Paragraph("3. Correcciones al informe de avance 1", styles["HeadingBlue"]))
    for item in [
        "El identificador se corrige a [a-zA-Z][a-zA-Z0-9_]* para que coincida con los ejemplos que contienen guion bajo.",
        "La cadena se define como comilla inicial, cero o más caracteres que no son comilla ni salto de línea y comilla final: CADENA : '\"' ~[\"\\r\\n]* '\"' ;. Se permiten cadenas vacías y no escapes.",
        "Se incorpora almacen como reservada ALMACEN.",
        "El signo menos queda separado de NUMERO.",
        "Los operadores de comparación, conjunción y disyunción tienen tokens separados. La nueva estructura de ExpresionLogica se implementará en el avance sintáctico.",
    ]:
        story.append(Paragraph("- " + item, styles["BodyES"]))

    lexical_table = pdf_table(
        [LEXICAL_HEADERS, *LEXICAL_ROWS],
        [3.2 * cm, 8.2 * cm, 5.6 * cm],
        styles,
    )
    story.append(KeepTogether([Paragraph("4. Componentes léxicos", styles["HeadingBlue"]), lexical_table]))

    story.append(Paragraph("5. Implementación con ANTLR 4", styles["HeadingBlue"]))
    story.append(Paragraph("DFDLexer.g4 contiene una regla por palabra reservada, operador, delimitador, literal y nombre. ANTLR elige la coincidencia más larga; en empates gana la regla anterior. Así, entidad es ENTIDAD, entidad1 es ID y -> es FLECHA.", styles["BodyES"]))
    story.append(Preformatted("ENTIDAD : 'entidad';\nFLECHA  : '->';\nNUMERO  : [0-9]+ ('.' [0-9]+)?;\nCADENA  : '\"' ~[\"\\r\\n]* '\"';\nID      : [a-zA-Z] [a-zA-Z0-9_]*;\nESPACIOS: [ \\t\\r\\n]+ -> skip;\nERROR   : .;", styles["CodeES"]))
    story.append(Paragraph("MostrarTokens.java lee un archivo UTF-8 y muestra TIPO | LEXEMA | LINEA | COLUMNA. El parser queda fuera de este avance.", styles["BodyES"]))
    story.append(Preformatted("mvn -q clean package\nmvn -q exec:java '-Dexec.mainClass=org.example.MostrarTokens' '-Dexec.args=src/test/resources/programa_informe1.dfd'", styles["CodeES"]))

    story.append(Paragraph("6. Pruebas del analizador", styles["HeadingBlue"]))
    story.append(Paragraph(f"El programa principal produjo {token_count('programa_informe1.txt')} tokens y ningún ERROR.", styles["BodyES"]))
    main_lines = lines("programa_informe1.txt")
    story.append(Preformatted("\n".join(main_lines[:8] + ["...", *main_lines[-6:]]), styles["CodeES"]))
    story.append(pdf_table([
        ["Entrada", "Resultado real resumido"],
        ["reservadas_identificadores.dfd", "Las 19 reservadas se distinguen de SISTEMA1, entidad1, cliente_externo y bd_01, que son ID."],
        ["numeros.dfd", "0, 12 y 3.5 son NUMERO; -2 es MENOS + NUMERO; puntos mal ubicados son ERROR."],
        ["cadenas.dfd", "Se reconocen \"Pedido\", \"\" y \"Cliente principal\" como CADENA."],
        ["operadores.dfd", "Los siete operadores de dos caracteres se reconocen individualmente."],
        ["invalidos.dfd", "@, #, $, y el guion bajo inicial son ERROR; 1proceso se divide en NUMERO e ID."],
        ["cadena_sin_cierre.dfd", "La comilla es ERROR y el texto posterior queda como ID. Limitación documentada."],
    ], [5.2 * cm, 11.8 * cm], styles))
    story.append(Paragraph("Las salidas completas se conservan en evidencia/salidas/. Las pruebas son léxicas y no validan sentencias completas.", styles["BodyES"]))

    story.append(Paragraph("7. Cronograma de la etapa", styles["HeadingBlue"]))
    story.append(pdf_table([
        ["Actividad", "Período", "Responsable", "Estado"],
        ["Revisar devolución y fijar lenguaje DFD", "Hasta 22/09/2026", "Equipo", "Realizado"],
        ["Ajustar vocabulario y lexer ANTLR", "22/09/2026", "Equipo", "Realizado"],
        ["Generar analizador y documentar pruebas", "22-23/09/2026", "Equipo", "Realizado con salidas guardadas"],
        ["Revisar datos del grupo, redacción y entrega", "23/09/2026", "Un integrante y equipo", "Verificación final pendiente"],
    ], [8.0 * cm, 3.0 * cm, 3.5 * cm, 2.5 * cm], styles))

    story.append(Paragraph("8. Trabajo previsto para el avance sintáctico", styles["HeadingBlue"]))
    story.append(Paragraph("En el avance 3 se reemplazará ExpresionLogica por niveles separados para comparación, conjunción y disyunción, manteniendo la precedencia de && sobre || y utilizando los tokens generados aquí. El parser no forma parte de esta entrega.", styles["BodyES"]))
    story.append(Paragraph("9. Archivos y verificaciones antes de entregar", styles["HeadingBlue"]))
    story.append(Paragraph("El archivo que corresponde presentar según el enunciado es el PDF. El DOCX se conserva como fuente editable y el proyecto con sus evidencias permite reproducir el analizador.", styles["BodyES"]))
    story.append(pdf_table([
        ["Elemento", "Acción"],
        [PDF_FILENAME, "Presentar después de completar los datos del grupo."],
        ["Informe_Avance_2_DFD.docx", "Conservar como editable; presentarlo solo si la cátedra lo solicita."],
        ["Proyecto Maven/ANTLR", "Conservar con pom.xml, DFDLexer.g4, MostrarTokens.java, pruebas y evidencia/salidas/."],
        ["Datos del grupo", "Completar legajos faltantes, confirmar nombres, grupo y continuidad de integrantes."],
    ], [5.0 * cm, 12.0 * cm], styles))
    story.append(Paragraph("Referencias", styles["HeadingBlue"]))
    story.append(Paragraph("Coronel, G. y Gili, J. (2026). Informe de avance 1: Taller 2026.<br/>Devolución de la entrega 1 - Coronel, Gili (2026).<br/>Cátedra de Lenguajes de Programación y Compiladores (2026). Enunciado del taller 2026. UNSE FCEyT.<br/>Las referencias legacy se consultaron solo para observar el formato histórico de la cátedra.", styles["BodyES"]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f"Informe de avance 2 - DFD - {doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
