from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "Informe_Avance_3_DFD.docx"
PDF_PATH = ROOT / "Informe_Avance_3_DFD.pdf"
FECHA = "6 de octubre de 2026"

GRAMATICA_RESUMIDA = """programa -> SISTEMA ID ; DefNivel* FIN_SISTEMA EOF
DefNivel -> NIVEL NUMERO Bloque
Bloque -> LLAVE_ABRE Sentencia* LLAVE_CIERRA
Sentencia -> DeclaracionNodo | DeclaracionFlujo | DeclaracionVar
           | Asignacion ; | Condicional | Iterativa | LlamadaFuncion ;
DeclaracionNodo -> TipoNodo ID = CADENA ;
TipoNodo -> ENTIDAD | PROCESO | ALMACEN
DeclaracionFlujo -> FLUJO ID = CONECTAR ( ID -> ID ) CON_DATO CADENA ;
DeclaracionVar -> Tipo ID = Expresion ;
Tipo -> ENTERO | REAL | BOOLEANO
Asignacion -> ID = Expresion
Condicional -> SI ( Expresion ) Bloque ( SINO Bloque )?
Iterativa -> MIENTRAS ( Expresion ) Bloque
LlamadaFuncion -> IMPRIMIR ( Expresion )
                | VALIDAR_MODELO ( NUMERO )
                | CONTAR_NODOS ( TipoNodo )

Expresion -> Disyuncion
Disyuncion -> Conjuncion ( OR Conjuncion )*
Conjuncion -> Negacion ( AND Negacion )*
Negacion -> NOT Negacion | Relacion
Relacion -> Suma ( OpRel Suma )?
OpRel -> == | != | < | > | <= | >=
Suma -> Producto ( (+ | -) Producto )*
Producto -> Unaria ( (* | /) Unaria )*
Unaria -> (+ | -) Unaria | Primaria
Primaria -> ID | NUMERO | CADENA | TRUE | FALSE
          | LlamadaFuncion | ( Expresion )"""

PRUEBAS = [
    (
        "programa_informe1.dfd",
        "Programa del informe anterior con sistema, nivel, nodos y flujo.",
        "Aceptado; 0 errores sintácticos.",
    ),
    (
        "sintaxis_precedencia_valida.dfd",
        "Tipos entero/real/booleano, operaciones unarias y binarias, not/&&/||, si/sino, mientras y funciones anidadas.",
        "Aceptado; 0 errores sintácticos.",
    ),
    (
        "sintaxis_comparacion_encadenada.dfd",
        "Comparación encadenada 1 < 2 < 3 con todos los lexemas reconocidos.",
        "Rechazado en línea 3, columna 21; 1 error sintáctico.",
    ),
    (
        "sintaxis_punto_y_coma_faltante.dfd",
        "Declaración sin punto y coma antes del cierre del bloque.",
        "Rechazado en línea 4, columna 0; 1 error sintáctico.",
    ),
]

CRONOGRAMA = [
    ("Revisión de la gramática corregida y alcance", FECHA, "Equipo Los Lexemas", "Realizado"),
    ("Ajuste de tokens léxicos requeridos por el parser", FECHA, "Equipo Los Lexemas", "Realizado"),
    ("Desarrollo de DFDParser.g4 con niveles de precedencia", FECHA, "Equipo Los Lexemas", "Realizado"),
    ("Pruebas de aceptación y rechazo con ANTLR TestRig", FECHA, "Equipo Los Lexemas", "Realizado"),
    ("Documentación y preparación del Informe de avance 3", FECHA, "Equipo Los Lexemas", "Realizado"),
]


def add_code_docx(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def add_reference_docx(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.first_line_indent = Cm(-0.7)
    for index, segment in enumerate(text.split("*")):
        run = paragraph.add_run(segment)
        run.italic = index % 2 == 1


def add_table_docx(document: Document, headers, rows) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    document.add_paragraph()


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for style_name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 11)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("1F4E79")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Informe de avance 3 - DFD Lang - Taller 2026")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Informe de avance Nº 3\nDesarrollo del analizador sintáctico con ANTLR 4")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Lenguajes de Programación y Compiladores - UNSE FCEyT - 2026")
    document.add_paragraph()
    document.add_paragraph("Lenguaje: DFD Lang")
    document.add_paragraph("Grupo: Los Lexemas")
    document.add_paragraph(
        "Integrantes: Gastón Coronel, Jerónimo Gili, Leonardo Noriega y Jorge Taboada."
    )
    document.add_paragraph(
        "Docentes: MSc. Ing. Margarita Álvarez y Lic. Eugenia Alejandra Gonzalez."
    )
    document.add_paragraph(f"Fecha: {FECHA}")
    document.add_page_break()

    document.add_heading("1. Objetivo y alcance", level=1)
    document.add_paragraph(
        "El objetivo de esta etapa es desarrollar con ANTLR 4 un analizador sintáctico "
        "para DFD Lang. El parser recibe los tokens producidos por DFDLexer y comprueba "
        "si la secuencia puede derivarse de la gramática del lenguaje. También permite "
        "observar el árbol de análisis y detectar errores sintácticos con su ubicación. "
        "El enunciado solicita el desarrollo y las pruebas de esta etapa "
        "(Cátedra de Lenguajes de Programación y Compiladores, 2026b, p. 2)."
    )
    document.add_paragraph(
        "El alcance comprende la estructura del programa, los niveles, bloques, "
        "declaraciones de nodos y flujos, variables, asignaciones, control de flujo, "
        "llamadas a funciones y expresiones. La comprobación de tipos, la existencia "
        "de identificadores y las reglas de conexión entre elementos DFD pertenecen "
        "a la etapa semántica y no se consideran validadas aquí."
    )

    document.add_heading("2. Fundamento teórico", level=1)
    document.add_paragraph(
        "Una gramática libre de contexto especifica las construcciones sintácticas "
        "válidas mediante símbolos terminales, no terminales y producciones. El "
        "analizador sintáctico recibe el flujo de tokens del analizador léxico y "
        "determina si puede generarse a partir de la gramática; además, informa "
        "errores sintácticos (Cátedra de Lenguajes de Programación y Compiladores, "
        "2026a, diap. 4-5). ANTLR genera el parser y el árbol de análisis a partir "
        "de las reglas declaradas en un archivo .g4 (ANTLR, s. f.)."
    )
    document.add_paragraph(
        "La gramática organiza las expresiones en niveles. Esta estructura representa "
        "la precedencia sin depender de una única producción ambigua: primero se "
        "agrupan las operaciones unarias y multiplicativas, luego suma/resta, "
        "relaciones, negación, conjunción y disyunción."
    )

    document.add_heading("3. Gramática implementada", level=1)
    document.add_paragraph(
        "El símbolo inicial es programa. El operador * representa repetición de cero "
        "o más elementos y ? indica una parte opcional. EOF exige consumir toda la "
        "entrada. La estructura de expresiones sigue la versión corregida del "
        "Informe de avance 2 (Coronel et al., 2026). Las reglas completas se encuentran en "
        "src/main/antlr4/org/example/DFDParser.g4."
    )
    add_code_docx(document, GRAMATICA_RESUMIDA)
    document.add_paragraph(
        "LLAVE_ABRE y LLAVE_CIERRA representan las llaves literales del bloque; "
        "Sentencia* indica cero o más sentencias. "
        "La precedencia implementada, de menor a mayor, es: OR, AND, NOT/relación, "
        "suma y resta, producto y división, operadores unarios y primarias. La regla "
        "Relacion admite como máximo un operador relacional; por ello, una expresión "
        "como 1 < 2 < 3 se rechaza. El signo de un número se analiza como operador "
        "unario separado del token NUMERO."
    )

    document.add_heading("4. Integración con ANTLR", level=1)
    document.add_paragraph(
        "DFDParser.g4 es una gramática de parser independiente y declara "
        "tokenVocab=DFDLexer para reutilizar los tokens existentes. En DFDLexer.g4 "
        "se incorporaron REAL, NOT y CADENA_SIN_CIERRE, necesarios para mantener "
        "la especificación corregida del Informe 2. Los tokens AND y OR ya estaban "
        "definidos. Maven ejecuta el plugin ANTLR 4.13.2 y compila las clases "
        "generadas. No se agregaron ni modificaron fuentes Java del proyecto."
    )
    add_code_docx(
        document,
        "mvn -q clean package\n"
        "./tools/verify_parser.ps1\n"
        "mvn -q exec:java '-Dexec.classpathScope=test' "
        "'-Dexec.mainClass=org.antlr.v4.gui.TestRig' "
        "'-Dexec.args=org.example.DFD programa -tree src/test/resources/programa_informe1.dfd'",
    )
    document.add_paragraph(
        "El árbol generado para 2 + 3 * 4 ubica 3 * 4 dentro del producto que forma "
        "el segundo operando de la suma. Para not x < 10 && true || false, el árbol "
        "ubica la relación dentro de la negación, la conjunción antes de la "
        "disyunción, de acuerdo con la precedencia especificada."
    )

    document.add_heading("5. Pruebas del analizador sintáctico", level=1)
    document.add_paragraph(
        "Se utilizó ANTLR TestRig sobre archivos de prueba. Los casos inválidos "
        "contienen tokens léxicamente reconocibles para comprobar que el diagnóstico "
        "corresponde a la fase sintáctica."
    )
    add_table_docx(document, ["Entrada", "Propósito", "Resultado"], PRUEBAS)
    document.add_paragraph("Diagnósticos obtenidos para los casos negativos:")
    add_code_docx(
        document,
        "line 3:21 mismatched input '<' expecting {'&&', '||', '+', '-', '*', '/', ';'}\n"
        "line 4:0 missing ';' at '}'",
    )
    document.add_paragraph(
        "El resumen reproducible se conserva en evidencia/salidas/sintaxis.txt y "
        "puede volver a generarse ejecutando tools/verify_parser.ps1."
    )

    document.add_heading("6. Cronograma de la etapa", level=1)
    add_table_docx(document, ["Actividad", "Fecha", "Responsable", "Estado"], CRONOGRAMA)

    document.add_heading("7. Conclusión", level=1)
    document.add_paragraph(
        "La gramática permite analizar la estructura principal de DFD Lang y separar "
        "la precedencia de operadores en reglas comprensibles. Las pruebas muestran "
        "la aceptación del programa de referencia y de expresiones con los "
        "constructos incorporados, además del rechazo de comparaciones encadenadas "
        "y sentencias incompletas. ANTLR facilita la generación del parser y del "
        "árbol; las validaciones semánticas quedan como trabajo de la siguiente etapa."
    )

    document.add_heading("Referencias", level=1)
    references = [
        "Cátedra de Lenguajes de Programación y Compiladores, UNSE-FCEyT. (2026a). "
        "*A. sintáctico* [Diapositivas de cátedra].",
        "Cátedra de Lenguajes de Programación y Compiladores, UNSE-FCEyT. (2026b). "
        "*Enunciado del taller 2026* [Consigna].",
        "Coronel, G., Gili, J., Noriega, L., & Taboada, J. (2026). "
        "*Informe de avance 2 corregido* [Documento de trabajo no publicado].",
        "ANTLR. (s. f.). *ANTLR 4*. https://www.antlr.org/",
    ]
    for reference in references:
        add_reference_docx(document, reference)

    document.save(DOCX_PATH)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="TitleBlue",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=21,
            textColor=colors.HexColor("#1F4E79"),
            alignment=TA_CENTER,
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeadingBlue",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#1F4E79"),
            spaceBefore=9,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyES",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.8,
            leading=11.5,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SmallES",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=9,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CodeES",
            parent=styles["Code"],
            fontName="Courier",
            fontSize=6.6,
            leading=7.8,
            leftIndent=7,
            rightIndent=7,
            spaceBefore=3,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReferenceES",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.2,
            leading=10.5,
            leftIndent=12,
            firstLineIndent=-12,
            spaceAfter=5,
        )
    )
    return styles


def pdf_table(rows, widths, styles):
    converted = [
        [Paragraph(escape(str(cell)), styles["SmallES"]) for cell in row]
        for row in rows
    ]
    table = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E8F5")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9FBAD0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def pdf_paragraph(text, style):
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def pdf_reference(text, style):
    segments = text.split("*")
    markup = "".join(
        f"<i>{escape(segment)}</i>" if index % 2 else escape(segment)
        for index, segment in enumerate(segments)
    )
    return Paragraph(markup, style)


def build_pdf() -> None:
    styles = pdf_styles()
    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=1.8 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    story = [
        Spacer(1, 2.5 * cm),
        pdf_paragraph("Informe de avance Nº 3", styles["TitleBlue"]),
        pdf_paragraph("Desarrollo del analizador sintáctico con ANTLR 4", styles["TitleBlue"]),
        pdf_paragraph("Lenguajes de Programación y Compiladores - UNSE FCEyT - 2026", styles["BodyES"]),
        Spacer(1, 1.2 * cm),
        pdf_paragraph("Lenguaje: DFD Lang", styles["BodyES"]),
        pdf_paragraph("Grupo: Los Lexemas", styles["BodyES"]),
        pdf_paragraph("Integrantes: Gastón Coronel, Jerónimo Gili, Leonardo Noriega y Jorge Taboada.", styles["BodyES"]),
        pdf_paragraph("Docentes: MSc. Ing. Margarita Álvarez y Lic. Eugenia Alejandra Gonzalez.", styles["BodyES"]),
        pdf_paragraph(f"Fecha: {FECHA}", styles["BodyES"]),
        PageBreak(),
    ]

    story.extend(
        [
            pdf_paragraph("1. Objetivo y alcance", styles["HeadingBlue"]),
            pdf_paragraph(
                "Desarrollar con ANTLR 4 un analizador sintáctico para DFD Lang. El parser recibe los tokens de DFDLexer y comprueba si la secuencia puede derivarse de la gramática; además, permite observar el árbol de análisis e informa errores sintácticos con ubicación. El enunciado solicita el desarrollo y las pruebas de esta etapa (Cátedra de Lenguajes de Programación y Compiladores, 2026b, p. 2).",
                styles["BodyES"],
            ),
            pdf_paragraph(
                "El alcance incluye programas, niveles, bloques, declaraciones DFD, variables, asignaciones, estructuras de control, funciones y expresiones. La comprobación de tipos, identificadores y conexiones corresponde a la etapa semántica y no se valida en este avance.",
                styles["BodyES"],
            ),
            pdf_paragraph("2. Fundamento teórico", styles["HeadingBlue"]),
            pdf_paragraph(
                "Una gramática libre de contexto describe las construcciones sintácticas mediante terminales, no terminales y producciones. El analizador sintáctico recibe el flujo de tokens del analizador léxico y determina si puede generarse a partir de la gramática (Cátedra de Lenguajes de Programación y Compiladores, 2026a, diap. 4-5). ANTLR genera el parser y el árbol de análisis desde reglas .g4 (ANTLR, s. f.).",
                styles["BodyES"],
            ),
            pdf_paragraph("3. Gramática implementada", styles["HeadingBlue"]),
            pdf_paragraph(
                "El símbolo inicial es programa. En la notación resumida, * representa cero o más repeticiones y ? una parte opcional. EOF exige consumir toda la entrada. La estructura de expresiones sigue la versión corregida del Informe de avance 2 (Coronel et al., 2026). Las reglas completas están en src/main/antlr4/org/example/DFDParser.g4.",
                styles["BodyES"],
            ),
            Preformatted(GRAMATICA_RESUMIDA, styles["CodeES"]),
            pdf_paragraph(
                "LLAVE_ABRE y LLAVE_CIERRA representan las llaves literales del bloque; Sentencia* indica cero o más sentencias. Precedencia de menor a mayor: OR, AND, NOT/relación, suma/resta, producto/división, unarios y primarias. Relacion admite como máximo un operador relacional, por lo que 1 < 2 < 3 se rechaza. Los signos numéricos se analizan como operadores unarios separados de NUMERO.",
                styles["BodyES"],
            ),
            pdf_paragraph("4. Integración con ANTLR", styles["HeadingBlue"]),
            pdf_paragraph(
                "DFDParser.g4 utiliza tokenVocab=DFDLexer para reutilizar los tokens existentes. DFDLexer.g4 incorpora REAL, NOT y CADENA_SIN_CIERRE; AND y OR ya estaban definidos. Maven utiliza ANTLR 4.13.2. No se modificaron fuentes Java del proyecto.",
                styles["BodyES"],
            ),
            Preformatted(
                "mvn -q clean package\n./tools/verify_parser.ps1\n"
                "TestRig: org.example.DFD programa -tree <archivo.dfd>",
                styles["CodeES"],
            ),
            pdf_paragraph(
                "El árbol de 2 + 3 * 4 agrupa el producto antes de la suma. En not x < 10 && true || false, la relación queda dentro de NOT, AND agrupa antes de OR y se conserva la precedencia definida.",
                styles["BodyES"],
            ),
            pdf_paragraph("5. Pruebas del analizador sintáctico", styles["HeadingBlue"]),
            pdf_paragraph(
                "Las entradas inválidas contienen lexemas reconocibles para que el diagnóstico corresponda a la fase sintáctica.",
                styles["BodyES"],
            ),
        ]
    )
    story.append(
        KeepTogether(
            [
                pdf_table(
                    [["Entrada", "Propósito", "Resultado"], *PRUEBAS],
                    [5.0 * cm, 6.3 * cm, 6.0 * cm],
                    styles,
                )
            ]
        )
    )
    story.extend(
        [
            Spacer(1, 5),
            pdf_paragraph("Diagnósticos de los casos negativos:", styles["BodyES"]),
            Preformatted(
                "line 3:21 mismatched input '<' expecting {'&&', '||', '+', '-', '*', '/', ';'}\n"
                "line 4:0 missing ';' at '}'",
                styles["CodeES"],
            ),
            pdf_paragraph(
                "El resumen de ejecución se conserva en evidencia/salidas/sintaxis.txt y se reproduce con tools/verify_parser.ps1.",
                styles["BodyES"],
            ),
            pdf_paragraph("6. Cronograma de la etapa", styles["HeadingBlue"]),
            pdf_table(
                [["Actividad", "Fecha", "Responsable", "Estado"], *CRONOGRAMA],
                [7.0 * cm, 2.5 * cm, 4.0 * cm, 3.8 * cm],
                styles,
            ),
            pdf_paragraph("7. Conclusión", styles["HeadingBlue"]),
            pdf_paragraph(
                "Las reglas implementadas validan la estructura principal de DFD Lang y expresan la precedencia de operadores en niveles separados. Las pruebas aceptan el programa de referencia y expresiones corregidas, y rechazan comparaciones encadenadas y sentencias incompletas. Las validaciones semánticas quedan para la siguiente etapa.",
                styles["BodyES"],
            ),
            pdf_paragraph("Referencias", styles["HeadingBlue"]),
        ]
    )

    references = [
        "Cátedra de Lenguajes de Programación y Compiladores, UNSE-FCEyT. (2026a). *A. sintáctico* [Diapositivas de cátedra].",
        "Cátedra de Lenguajes de Programación y Compiladores, UNSE-FCEyT. (2026b). *Enunciado del taller 2026* [Consigna].",
        "Coronel, G., Gili, J., Noriega, L., & Taboada, J. (2026). *Informe de avance 2 corregido* [Documento de trabajo no publicado].",
        "ANTLR. (s. f.). *ANTLR 4*. https://www.antlr.org/",
    ]
    story.extend(pdf_reference(item, styles["ReferenceES"]) for item in references)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawCentredString(
            A4[0] / 2,
            0.8 * cm,
            f"Informe de avance 3 - DFD Lang - {doc.page}",
        )
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
